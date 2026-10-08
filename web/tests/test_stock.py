import os
import tempfile
import unittest

from web import stock


def make_workbook(path):
    """A small copy of the real layouts: UK-style (3 header rows, merged groups) and an
    older 2-header-row tab, with subtotal, grand-total and legend rows."""
    import openpyxl
    wb = openpyxl.Workbook()
    uk = wb.active
    uk.title = "UK"
    uk["B1"] = "UK"
    uk["F1"] = "ON HAND STOCK"
    uk.merge_cells("F1:M1")
    uk["N1"] = "INBOUND UNITS FROM CHINA"
    uk.merge_cells("N1:T1")
    uk["U1"] = "ONHAND + INBOUND STOCK"
    uk.merge_cells("U1:W1")
    uk["X1"] = "Remarks"
    for col, name in zip("ABCDEFGHIJKLM", ["Products", "SKU", "ASIN", "7 days", "30 days", "Units in Amazon",
                                           "AMZ days stock lasts (7)", "AMZ days stock lasts (30)",
                                           "Transit to AMZ from 3 PL", "Stocks with 3PL", "TOTAL UNITS",
                                           "Days stock last (7 days)", "Days stock last (30 days)"]):
        uk[f"{col}2"] = name
    uk["N2"] = "Transit to Amazon"
    uk.merge_cells("N2:P2")
    uk["N3"], uk["O3"], uk["P3"] = "AIR", "FAST OCEAN", "OCEAN"
    uk["Q2"], uk["R2"], uk["S2"] = "ETA (Air)", "TOTAL UNITS", "Days stock last (7 days)"
    uk["U2"], uk["V2"] = "Total units", "Days stock last (7 days)"
    rows = [
        ["Pill Crusher V.2", "PMPC-UK002", "B075HZYT9G", 4.43, 4.83, 685, 155, 142, 0, 0, 685, 155, 142,
         0, 0, 50, None, 50, 11, None, 735, 166, None, "ok"],
        ["Electric Pill Crusher", "PMEC-001", "B093QFLW1T", 1.14, 0.97, 60, 53, 62, 0, 0, 60, 53, 62,
         0, 0, 100, None, 100, 88, None, 160, 140, None, "SEND STOCK TO AMAZON"],
        ["Total Pill Crusher", None, None, 5.57, 5.8],
        ["Ice Pack", "TCIP-001", "B0X", 0.01, 0.01, 500, 0, 0, 0, 0, 500],
        ["Gloves", "GL-1", "B0Y", "#N/A", "-", 10],
        ["Chemo Cap", "CC-1", "B0Z", "12", "10", "0", None, None, None, None, "0", None, None],
        [None, None, None, None, None, 20288],
        ["LEGENDS:"],
    ]
    for i, r in enumerate(rows, start=4):
        for j, v in enumerate(r, start=1):
            uk.cell(i, j, v)
    old = wb.create_sheet("USA old")
    for j, name in enumerate(["Products", "SKU", "ASIN", "Average Sales (7 days)", "Average Sales (30 days)",
                              "Units in Amazon"], start=1):
        old.cell(2, j, name)
    for j, v in enumerate(["Widget", "W-1", "B01", 10, 20, 300], start=1):
        old.cell(3, j, v)
    notes = wb.create_sheet("SKU Status")
    notes.append(["Products", "SKU", "ASIN", "Market"])
    wb.save(path)


class StockTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fd, cls.path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        make_workbook(cls.path)

    @classmethod
    def tearDownClass(cls):
        os.remove(cls.path)

    def test_tab_names(self):
        self.assertEqual(stock.tab_names(self.path), [("UK", True), ("USA old", True), ("SKU Status", False)])

    def test_columns_and_rows(self):
        uk, old = stock.read_tabs(self.path, ["UK", "USA old"])
        self.assertEqual(uk.problem, "")
        self.assertEqual(uk.columns["onhand"], 10)    # K: first TOTAL UNITS, under ON HAND STOCK
        self.assertEqual(uk.columns["inbound"], 17)   # R: TOTAL UNITS under INBOUND
        self.assertEqual(uk.columns["total"], 20)     # U: Total units under ONHAND + INBOUND
        self.assertEqual([i.sku for i in uk.items], ["PMPC-UK002", "PMEC-001", "TCIP-001", "GL-1", "CC-1"])
        pm = uk.items[1]
        self.assertEqual((pm.sales7, pm.sales30, pm.amazon, pm.onhand, pm.inbound, pm.total),
                         (1.14, 0.97, 60, 60, 100, 160))
        self.assertEqual(pm.remarks, "SEND STOCK TO AMAZON")
        self.assertEqual(old.items[0].onhand, 300)    # no on-hand total: Amazon units
        self.assertEqual(old.items[0].total, 300)

    def test_days_and_levels(self):
        tabs = stock.read_tabs(self.path, ["UK", "USA old"])
        rows, counts = stock.risk_rows(tabs, "onhand", "max", 60)
        by = {r["item"].sku: r for r in rows}
        self.assertAlmostEqual(by["PMEC-001"]["days"], 60 / 1.14)          # higher of 7/30-day sales
        self.assertEqual(by["PMEC-001"]["level"], "low")
        self.assertAlmostEqual(by["PMPC-UK002"]["days"], 685 / 4.83)
        self.assertEqual(by["PMPC-UK002"]["level"], "ok")
        self.assertEqual(by["TCIP-001"]["level"], "nosales")                # 0.01 = no real sales
        self.assertEqual(by["GL-1"]["level"], "nosales")
        self.assertEqual(by["CC-1"]["level"], "out")
        self.assertEqual(by["W-1"]["days"], 15)
        self.assertEqual(by["W-1"]["level"], "critical")
        self.assertEqual(rows[0]["item"].sku, "CC-1")                       # most urgent first
        self.assertEqual(counts, {"out": 1, "critical": 1, "low": 1, "ok": 1, "nosales": 2})
        rows, _ = stock.risk_rows(tabs, "inbound", "30", 60)
        self.assertEqual({r["item"].sku: r["level"] for r in rows}["PMEC-001"], "ok")   # 160 / 0.97


if __name__ == "__main__":
    unittest.main()
