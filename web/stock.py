"""Read the stock planning sheet (Google Sheet rows or an .xlsx) and work out days of stock per SKU.

Each market tab has 2-3 header rows: row 1 = group ("ON HAND STOCK", "INBOUND UNITS FROM
CHINA", "ONHAND + INBOUND STOCK", ...), row 2 = column name, row 3 (newer tabs) = route
(AIR / FAST OCEAN / OCEAN). Layouts differ per tab, so columns are found by their header
text, not by letter. Sales columns are average units per day (7 and 30 days).
Rows without a SKU (family subtotals, grand totals, legends) are skipped.
"""
import re
from dataclasses import dataclass, field

NO_SALES = 0.01   # the sheet floors sales at 0.01, which would give absurd days of stock

STOCK_MEASURES = {
    "amazon": "Amazon only",
    "onhand": "On hand (Amazon + AWD + 3PL)",
    "inbound": "On hand + inbound",
}
SALES_MEASURES = {
    "max": "Higher of 7-day and 30-day sales",
    "7": "7-day sales",
    "30": "30-day sales",
}


@dataclass
class Item:
    tab: str
    row: int
    sku: str
    asin: str = ""
    product: str = ""
    sales7: float = None
    sales30: float = None
    amazon: float = None
    onhand: float = None
    inbound: float = None
    total: float = None            # on hand + inbound
    remarks: str = ""

    def sales(self, which):
        vals = {"7": [self.sales7], "30": [self.sales30], "max": [self.sales7, self.sales30]}[which]
        vals = [v for v in vals if v is not None]
        return max(vals) if vals else None

    def units(self, which):
        return {"amazon": self.amazon, "onhand": self.onhand, "inbound": self.total}[which]

    def days(self, stock, sales):
        """Days the stock lasts, or None when there are no sales to divide by."""
        rate, units = self.sales(sales), self.units(stock)
        if rate is None or rate <= NO_SALES or units is None:
            return None
        return max(units, 0) / rate


@dataclass
class Tab:
    name: str
    columns: dict = field(default_factory=dict)   # role -> column index (0-based)
    items: list = field(default_factory=list)
    problem: str = ""


def norm(v):
    return re.sub(r"\s+", " ", str(v or "")).strip().lower()


def number(v):
    """Cell value -> float, or None for blanks, dashes, '#N/A' and other text."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip().replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", s):
        return float(s)
    return None


def header_labels(rows):
    """rows: the header rows (2 or 3 tuples). Returns per column (group, name, sub), with
    merged group/name cells filled to the right the way they appear on screen."""
    width = max(len(r) for r in rows)
    get = lambda r, i: rows[r][i] if r < len(rows) and i < len(rows[r]) else None
    labels, group, name = [], "", ""
    for i in range(width):
        g, n, s = norm(get(0, i)), norm(get(1, i)), norm(get(2, i)) if len(rows) > 2 else ""
        group = g or group
        if n:
            name = n
        elif not s:
            name = ""        # a name only spreads over the route columns under it
        labels.append((group, name if (n or s) else "", s))
    return labels


SALES7 = re.compile(r"^(average sales )?\(?7 days\)?$")
SALES30 = re.compile(r"^(average sales )?\(?30 days\)?$")


def find_columns(labels):
    """Map the columns this report needs to their index. First match wins, because the
    sheet repeats some names (Walmart sales come after Amazon sales)."""
    cols = {}

    def take(role, test):
        if role in cols:
            return
        for i, (g, n, s) in enumerate(labels):
            if test(g, n, s):
                cols[role] = i
                return

    take("product", lambda g, n, s: n in ("products", "product", "desc", "products name"))
    take("sku", lambda g, n, s: n == "sku")
    take("asin", lambda g, n, s: n == "asin")
    take("sales7", lambda g, n, s: bool(SALES7.match(n)))
    take("sales30", lambda g, n, s: bool(SALES30.match(n)))
    take("amazon", lambda g, n, s: n == "units in amazon")
    take("total", lambda g, n, s: "onhand + inbound" in g.replace("on hand", "onhand") and n == "total units")
    take("inbound", lambda g, n, s: "inbound" in g and "onhand" not in g.replace("on hand", "onhand")
         and n == "total units")
    take("onhand", lambda g, n, s: ("on hand" in g or "onhand" in g) and "inbound" not in g and n == "total units")
    take("remarks", lambda g, n, s: g == "remarks" or n == "remarks")
    if "product" not in cols and "sku" in cols and cols["sku"] > 0:
        cols["product"] = 0
    return cols


def header_rows(rows):
    """2 or 3: newer tabs have a third header row (AIR / FAST OCEAN / OCEAN) with no SKU in it;
    on older tabs row 3 is already the first product."""
    if len(rows) < 3:
        return 2
    cols = find_columns(header_labels(rows[:2]))
    third = rows[2]
    sku = third[cols["sku"]] if "sku" in cols and cols["sku"] < len(third) else None
    return 2 if sku not in (None, "") else 3


def looks_like_stock(top_rows):
    """True if a tab's first 3 rows have SKU, sales and Amazon stock headers."""
    if not top_rows:
        return False
    cols = find_columns(header_labels(top_rows[:header_rows(top_rows)]))
    return all(k in cols for k in ("sku", "sales7", "amazon"))


def read_tab(ws, name):
    return parse_rows(list(ws.iter_rows(values_only=True)), name)


def parse_rows(rows, name):
    """rows: the whole tab as lists of cell values, top row first."""
    tab = Tab(name)
    if len(rows) < 3:
        tab.problem = "too few rows"
        return tab
    start = header_rows(rows)
    labels = header_labels(rows[:start])
    cols = find_columns(labels)
    tab.columns = cols
    missing = [r for r in ("sku", "sales7", "sales30", "amazon") if r not in cols]
    if missing:
        tab.problem = "can't find column(s): " + ", ".join(missing)
        return tab
    cell = lambda r, role: r[cols[role]] if role in cols and cols[role] < len(r) else None
    for i, r in enumerate(rows[start:], start=start + 1):
        sku = str(cell(r, "sku") or "").strip()
        product = str(cell(r, "product") or "").strip()
        if not sku or sku.startswith("#") or product.lower().startswith("total"):
            continue
        it = Item(tab=name, row=i, sku=sku, asin=str(cell(r, "asin") or "").strip(), product=product,
                  sales7=number(cell(r, "sales7")), sales30=number(cell(r, "sales30")),
                  amazon=number(cell(r, "amazon")), onhand=number(cell(r, "onhand")),
                  inbound=number(cell(r, "inbound")), total=number(cell(r, "total")),
                  remarks=str(cell(r, "remarks") or "").strip())
        if it.onhand is None:
            it.onhand = it.amazon
        if it.total is None and it.onhand is not None:
            it.total = it.onhand + (it.inbound or 0)
        tab.items.append(it)
    if not tab.items:
        tab.problem = "no SKU rows found"
    return tab


def open_workbook(path):
    import openpyxl
    # data_only: read the values Excel/Sheets last calculated, not the formulas
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def tab_names(path):
    """All tabs, each with whether it looks like a stock tab (has SKU and sales headers)."""
    wb = open_workbook(path)
    try:
        return [(ws.title, looks_like_stock(list(ws.iter_rows(min_row=1, max_row=3, values_only=True))))
                for ws in wb.worksheets]
    finally:
        wb.close()


def read_tabs(path, names):
    wb = open_workbook(path)
    try:
        have = set(wb.sheetnames)
        return [read_tab(wb[n], n) if n in have else Tab(n, problem="tab not in this file") for n in names]
    finally:
        wb.close()


def risk_rows(tabs, stock="onhand", sales="max", threshold=60):
    """Every SKU with its days of stock; returns (rows sorted most urgent first, counts)."""
    rows = []
    for tab in tabs:
        for it in tab.items:
            d = it.days(stock, sales)
            rate = it.sales(sales)
            if rate is None or rate <= NO_SALES:
                level = "nosales"
            elif d is None:
                level = "nostock"     # selling, but the stock cell is empty or text: someone should look
            elif d < 1:
                level = "out"
            elif d < min(30, threshold):
                level = "critical"
            elif d < threshold:
                level = "low"
            else:
                level = "ok"
            rows.append({"item": it, "days": d, "rate": rate, "level": level,
                         "days_amazon": it.days("amazon", sales), "days_onhand": it.days("onhand", sales),
                         "days_total": it.days("inbound", sales)})
    order = {"out": 0, "critical": 1, "low": 2, "nostock": 3, "ok": 4, "nosales": 5}
    rows.sort(key=lambda r: (order[r["level"]], r["days"] if r["days"] is not None else 1e9))
    counts = {k: sum(1 for r in rows if r["level"] == k) for k in order}
    return rows, counts
